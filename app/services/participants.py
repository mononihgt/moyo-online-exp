import csv
import io
import secrets
import zipfile

from fastapi import HTTPException, status
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from ..constants import (
    PARTICIPANT_COMPLETED,
    PARTICIPANT_DISABLED,
    PARTICIPANT_ELIGIBLE,
    VERSION_PUBLISHED,
)
from ..models import ConsentVersion, Participant, RoomParticipant
from ..security import decrypt_phone, encrypt_phone, mask_phone, normalize_phone, phone_hash
from ..timeutil import isoformat_utc

_NAME_HEADERS = {"name", "姓名"}
_PHONE_HEADERS = {"phone", "手机号", "手机", "phone number", "phonenumber"}
_NOTES_HEADERS = {"notes", "note", "备注", "说明"}


def _header_key(value) -> str:
    return str(value or "").strip().lower()


def parse_roster(data: bytes, filename: str) -> list[dict]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        text = data.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
    elif suffix in ("xlsx", "xlsm"):
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except (zipfile.BadZipFile, OSError) as exc:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "XLSX 文件无法读取") from exc
        sheet = workbook.active
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    else:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "名单仅接受 CSV 或 XLSX 文件")
    if not rows:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "名单为空")
    header = [_header_key(cell) for cell in rows[0]]
    name_index = next((index for index, value in enumerate(header) if value in _NAME_HEADERS), None)
    phone_index = next((index for index, value in enumerate(header) if value in _PHONE_HEADERS), None)
    notes_index = next((index for index, value in enumerate(header) if value in _NOTES_HEADERS), None)
    if name_index is None or phone_index is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "名单必须包含姓名和手机号列")
    result = []
    for row_number, row in enumerate(rows[1:], start=2):
        if row is None or all(value in (None, "") for value in row):
            continue
        name = str(row[name_index]).strip() if name_index < len(row) and row[name_index] is not None else ""
        phone_value = row[phone_index] if phone_index < len(row) else None
        if isinstance(phone_value, float) and phone_value.is_integer():
            phone_raw = str(int(phone_value))
        else:
            phone_raw = phone_value
        notes = None
        if notes_index is not None and notes_index < len(row) and row[notes_index] is not None:
            notes = str(row[notes_index]).strip() or None
        result.append({"row": row_number, "name": name, "phone": str(phone_raw or ""), "notes": notes})
    return result


async def import_completed(db: AsyncSession, data: bytes, filename: str) -> dict:
    rows = parse_roster(data, filename)
    result = {"total": len(rows), "created": 0, "matched": 0, "unmatched": 0, "rows": []}
    seen_hashes: set[str] = set()

    def record(row: dict, row_status: str, error: str | None) -> None:
        result["rows"].append(
            {
                "row": row["row"],
                "name": row["name"],
                "phoneMasked": _safe_mask(row["phone"]),
                "status": row_status,
                "error": error,
            }
        )

    for row in rows:
        name = row["name"].strip()
        if not name:
            result["unmatched"] += 1
            record(row, "unmatched", "姓名为空")
            continue
        if len(name) > 120:
            result["unmatched"] += 1
            record(row, "unmatched", "姓名过长")
            continue
        try:
            phone = normalize_phone(row["phone"])
        except ValueError as exc:
            result["unmatched"] += 1
            record(row, "unmatched", str(exc))
            continue
        digest = phone_hash(phone)
        if digest in seen_hashes:
            result["unmatched"] += 1
            record(row, "unmatched", "名单内手机号重复")
            continue
        seen_hashes.add(digest)
        query = await db.execute(select(Participant).where(Participant.phone_hash == digest))
        participant = query.scalars().first()
        if participant is None:
            db.add(
                Participant(
                    name=name,
                    phone_hash=digest,
                    phone_encrypted=encrypt_phone(phone),
                    status=PARTICIPANT_COMPLETED,
                    notes=row["notes"],
                )
            )
            await db.flush()
            result["created"] += 1
            record(row, "created", None)
            continue
        if participant.name.strip() != name:
            result["unmatched"] += 1
            record(row, "unmatched", "姓名与已有记录不一致")
            continue
        participant.status = PARTICIPANT_COMPLETED
        result["matched"] += 1
        record(row, "matched", None)
    await db.commit()
    return result


def _safe_mask(raw: str) -> str:
    try:
        return mask_phone(normalize_phone(raw))
    except ValueError:
        return "****"


async def login_or_register_participant(db: AsyncSession, name: str, phone_raw: str) -> Participant:
    try:
        phone = normalize_phone(phone_raw)
    except ValueError as exc:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "姓名或手机号不正确") from exc
    cleaned_name = name.strip()
    if not cleaned_name:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "姓名或手机号不正确")
    digest = phone_hash(phone)

    async def find_existing() -> Participant | None:
        query = await db.execute(select(Participant).where(Participant.phone_hash == digest))
        return query.scalars().first()

    participant = await find_existing()
    if participant is not None:
        if participant.name.strip() != cleaned_name:
            raise HTTPException(status.HTTP_401_UNAUTHORIZED, "姓名或手机号不正确")
        return participant

    participant = Participant(
        name=cleaned_name,
        phone_hash=digest,
        phone_encrypted=encrypt_phone(phone),
        status=PARTICIPANT_ELIGIBLE,
    )
    db.add(participant)
    try:
        await db.commit()
    except IntegrityError as exc:
        await db.rollback()
        participant = await find_existing()
        if participant is None or participant.name.strip() != cleaned_name:
            raise HTTPException(status.HTTP_401_UNAUTHORIZED, "姓名或手机号不正确") from exc
        return participant
    return participant


async def current_consent_version(db: AsyncSession) -> ConsentVersion | None:
    query = await db.execute(
        select(ConsentVersion)
        .where(ConsentVersion.status == VERSION_PUBLISHED)
        .order_by(ConsentVersion.version.desc())
        .limit(1)
    )
    return query.scalars().first()


async def participant_me(db: AsyncSession, participant: Participant) -> dict:
    from .rooms import one_time_experiment_lock

    consent = await current_consent_version(db)
    consent_required = consent is not None and participant.consent_version_id != consent.id
    completed_count_query = await db.execute(
        select(func.count(RoomParticipant.id)).where(
            RoomParticipant.participant_id == participant.id, RoomParticipant.completed_at.is_not(None)
        )
    )
    completed_count = completed_count_query.scalar_one()
    lock = await one_time_experiment_lock(db, participant.id)
    return {
        "id": participant.id,
        "name": participant.name,
        "phoneMasked": mask_phone(decrypt_phone(participant.phone_encrypted)),
        "status": participant.status,
        "canBook": participant.status == PARTICIPANT_ELIGIBLE and not consent_required and not lock.locked,
        "hasPendingQuestionnaire": lock.pending_questionnaire,
        "hasPayment": lock.has_payment,
        "consentRequired": consent_required,
        "consentVersion": (
            {"id": consent.id, "version": consent.version, "publishedAt": isoformat_utc(consent.published_at)}
            if consent is not None
            else None
        ),
        "completedExperiments": completed_count,
        "isDisabled": participant.status == PARTICIPANT_DISABLED,
        "isCompleted": participant.status == PARTICIPANT_COMPLETED,
    }


def new_analysis_code() -> str:
    return "P-" + secrets.token_hex(6)
